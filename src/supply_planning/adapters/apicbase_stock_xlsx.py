from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from supply_planning.adapters.errors import InputFileError
from supply_planning.domain.models import (
    InventorySnapshot,
    Item,
    ItemPlanningPolicy,
    Provenance,
    StockQuantityUnit,
)

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - exercised only in a broken installation
    load_workbook = None  # type: ignore[assignment]


STOCK_HEADERS = (
    "Stock Item Name",
    "UID",
    "Acc. Category",
    "Category",
    "Subcategory",
    "Supplier Article Numbers",
    "Current Stock (qty)",
    "Current Stock (value)",
    "Par",
    "Diff. with Par",
    "Min. Stock",
    "Diff with Min. Stock",
)
_REPORT_LOCATION_RE = re.compile(r"^Stock Report for\s+(.+?)\s*$", re.IGNORECASE)
_EXPORTED_AT_RE = re.compile(r"\bam\s+(\d{2}\.\d{2}\.\d{4})\s+(\d{2}:\d{2})\b")


@dataclass(frozen=True, slots=True)
class StockMappingReview:
    item_id: str
    item_name: str
    apicbase_uid: str | None
    apicbase_stock_item_name: str | None
    reason: str
    remedy: str

    def as_dict(self) -> dict[str, str]:
        return {
            "item_id": self.item_id,
            "item_name": self.item_name,
            "apicbase_uid": self.apicbase_uid or "",
            "apicbase_stock_item_name": self.apicbase_stock_item_name or "",
            "reason": self.reason,
            "remedy": self.remedy,
        }


@dataclass(frozen=True, slots=True)
class StockNormalizationResult:
    snapshots: tuple[InventorySnapshot, ...]
    reviews: tuple[StockMappingReview, ...]
    counted_at: datetime
    source_report_location: str
    source_version: str
    content_hash: str


def _content_hash(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
    except OSError as exc:
        raise InputFileError(f"cannot read {path}: {exc}") from exc
    return digest.hexdigest()


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _quantity(path: Path, row_number: int, value: Any) -> Decimal:
    try:
        quantity = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise InputFileError(
            f"{path} sheet Stock Report row {row_number}, field Current Stock (qty): "
            f"{value!r} is not numeric"
        ) from exc
    if quantity < 0:
        raise InputFileError(
            f"{path} sheet Stock Report row {row_number}, field Current Stock (qty): "
            "negative stock is not supported; correct the export or confirm adjustment semantics"
        )
    return quantity


def normalize_apicbase_stock(
    path: Path,
    *,
    location_id: str,
    timezone_name: str,
    items: Iterable[Item],
    item_policies: Iterable[ItemPlanningPolicy],
    required_item_ids: Iterable[str],
    assume_zero_for_unmapped: bool = False,
) -> StockNormalizationResult:
    if load_workbook is None:
        raise InputFileError(
            "XLSX support is unavailable; install the project dependency openpyxl>=3.1,<4"
        )
    if not location_id.strip():
        raise InputFileError("selected location_id must not be blank")
    try:
        timezone = ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError as exc:
        raise InputFileError(f"unknown planning timezone {timezone_name!r}") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError) as exc:
        raise InputFileError(f"cannot read Apicbase stock XLSX {path}: {exc}") from exc
    try:
        if "Stock Report" not in workbook.sheetnames:
            raise InputFileError(
                f"{path}: missing sheet 'Stock Report'; export the current stock list from Apicbase"
            )
        sheet = workbook["Stock Report"]
        report_title = _text(sheet.cell(row=1, column=1).value)
        report_match = _REPORT_LOCATION_RE.match(report_title)
        if report_match is None:
            raise InputFileError(
                f"{path} sheet Stock Report cell A1: expected 'Stock Report for <location>', "
                f"found {report_title!r}"
            )
        source_report_location = report_match.group(1).strip()

        export_text = " ".join(
            _text(sheet.cell(row=1, column=column).value)
            for column in range(1, len(STOCK_HEADERS) + 1)
            if sheet.cell(row=1, column=column).value not in (None, "")
        )
        export_match = _EXPORTED_AT_RE.search(export_text)
        if export_match is None:
            raise InputFileError(
                f"{path} sheet Stock Report row 1: cannot find export timestamp like "
                "'am 26.08.2026 17:10'; re-export the report"
            )
        counted_at = datetime.strptime(
            f"{export_match.group(1)} {export_match.group(2)}", "%d.%m.%Y %H:%M"
        ).replace(tzinfo=timezone)

        supplied_headers = tuple(
            _text(sheet.cell(row=3, column=column).value)
            for column in range(1, len(STOCK_HEADERS) + 1)
        )
        if supplied_headers != STOCK_HEADERS:
            raise InputFileError(
                f"{path} sheet Stock Report row 3: unsupported Apicbase export columns; "
                "download the standard current Stock Report without changing its headers"
            )

        source_rows: list[dict[str, Any]] = []
        for row_number in range(4, sheet.max_row + 1):
            values = [
                sheet.cell(row=row_number, column=column).value
                for column in range(1, len(STOCK_HEADERS) + 1)
            ]
            if not any(value not in (None, "") for value in values):
                continue
            row = dict(zip(STOCK_HEADERS, values, strict=True))
            name = _text(row["Stock Item Name"])
            if not name:
                raise InputFileError(
                    f"{path} sheet Stock Report row {row_number}: Stock Item Name is blank"
                )
            row["_row_number"] = row_number
            row["_name"] = name
            row["_uid"] = _text(row["UID"])
            source_rows.append(row)
    finally:
        workbook.close()

    by_uid: dict[str, dict[str, Any]] = {}
    by_name: dict[str, dict[str, Any]] = {}
    for row in source_rows:
        uid = row["_uid"]
        name = row["_name"]
        if uid:
            if uid in by_uid:
                raise InputFileError(
                    f"{path} sheet Stock Report: UID {uid!r} occurs on rows "
                    f"{by_uid[uid]['_row_number']} and {row['_row_number']}; resolve the duplicate"
                )
            by_uid[uid] = row
        if name in by_name:
            raise InputFileError(
                f"{path} sheet Stock Report: Stock Item Name {name!r} occurs on rows "
                f"{by_name[name]['_row_number']} and {row['_row_number']}; resolve the duplicate"
            )
        by_name[name] = row

    items_by_id = {item.item_id: item for item in items}
    policies_by_id = {policy.item_id: policy for policy in item_policies}
    snapshots: list[InventorySnapshot] = []
    reviews: list[StockMappingReview] = []
    for item_id in sorted(set(required_item_ids)):
        item = items_by_id.get(item_id)
        policy = policies_by_id.get(item_id)
        if item is None or policy is None:
            raise InputFileError(
                f"master data is missing item/policy for required BOM item_id={item_id!r}"
            )
        uid_match = by_uid.get(policy.apicbase_uid or "")
        name_match = by_name.get(policy.apicbase_stock_item_name or "")
        if uid_match is not None and name_match is not None and uid_match is not name_match:
            raise InputFileError(
                f"{path}: item_id={item_id!r} maps by UID to row {uid_match['_row_number']} "
                f"but by exact name to row {name_match['_row_number']}; correct the master mapping"
            )
        matched = uid_match or name_match
        if matched is None:
            reason = (
                "no reviewed Apicbase UID or exact stock-name match"
                if policy.apicbase_uid or policy.apicbase_stock_item_name
                else "item has no Apicbase UID or exact stock-name mapping"
            )
            reviews.append(
                StockMappingReview(
                    item_id=item_id,
                    item_name=item.item_name,
                    apicbase_uid=policy.apicbase_uid,
                    apicbase_stock_item_name=policy.apicbase_stock_item_name,
                    reason=reason,
                    remedy="Fill or correct apicbase_uid/apicbase_stock_item_name in Items and rerun.",
                )
            )
            if assume_zero_for_unmapped:
                snapshots.append(
                    InventorySnapshot(
                        location_id=location_id,
                        item_id=item_id,
                        counted_at=counted_at,
                        usable_on_hand_units=Decimal("0"),
                        partial_pack_g=Decimal("0"),
                        provenance=Provenance.POLICY_DEFAULT,
                    )
                )
            continue

        quantity = _quantity(path, int(matched["_row_number"]), matched["Current Stock (qty)"])
        usable_pack_units = (
            quantity
            if policy.stock_qty_unit is StockQuantityUnit.PACK
            else quantity * policy.packs_per_order_unit
        )
        snapshots.append(
            InventorySnapshot(
                location_id=location_id,
                item_id=item_id,
                counted_at=counted_at,
                usable_on_hand_units=usable_pack_units,
                partial_pack_g=Decimal("0"),
                provenance=Provenance.OBSERVED,
            )
        )

    content_hash = _content_hash(path)
    return StockNormalizationResult(
        snapshots=tuple(snapshots),
        reviews=tuple(reviews),
        counted_at=counted_at,
        source_report_location=source_report_location,
        source_version=f"apicbase-stock-{content_hash[:12]}",
        content_hash=content_hash,
    )
