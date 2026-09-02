from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from supply_planning.adapters.errors import InputFileError
from supply_planning.domain.models import (
    BomLine,
    DeliveryCoverageRule,
    ForecastDaily,
    Item,
    ItemPlanningPolicy,
    ItemType,
    Location,
    MenuCalendarEntry,
    Provenance,
    ShelfLifeAnchor,
    StockQuantityUnit,
    StorageClass,
)

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - exercised only in a broken installation
    load_workbook = None  # type: ignore[assignment]


ITEM_HEADERS = (
    "item_id",
    "item_name",
    "item_type",
    "official_supplier",
    "ordering_channel",
    "po_article_no",
    "po_description_match",
    "inner_pack_article_no",
    "inner_pack_ean",
    "order_unit_ean",
    "storage_class",
    "pack_size_g",
    "packs_per_order_unit",
    "order_unit",
    "stock_qty_unit",
    "apicbase_uid",
    "apicbase_stock_item_name",
    "lead_time_calendar_days",
    "shelf_life_days",
    "shelf_life_anchor",
    "min_safety_days",
    "yield_factor",
    "max_cover_days",
    "moq_order_units",
    "case_multiple_order_units",
    "active",
    "data_status",
    "source_note",
)
LOCATION_HEADERS = (
    "location_id",
    "location_name",
    "timezone",
    "active",
    "data_status",
    "source_note",
)
DELIVERY_HEADERS = (
    "delivery_rule_id",
    "location_id",
    "ordering_channel",
    "storage_class",
    "delivery_weekday",
    "covered_service_days",
    "order_weekday",
    "order_cutoff_local",
    "receipt_available_local",
    "review_period_days",
    "effective_from",
    "effective_to",
    "active",
    "data_status",
    "source_note",
)
DEMAND_HEADERS = (
    "location_id",
    "service_date",
    "dish_id",
    "dish_name",
    "forecast_portions",
    "forecast_version",
    "provenance",
    "source_note",
)
MENU_HEADERS = (
    "location_id",
    "service_date",
    "dish_id",
    "dish_name",
    "menu_version",
    "active",
    "provenance",
    "source_note",
)
BOM_HEADERS = (
    "bom_line_id",
    "dish_id",
    "dish_name",
    "silo_id",
    "silo_name",
    "item_id",
    "grams_per_portion",
    "effective_from",
    "effective_to",
    "bom_version",
    "active",
    "data_status",
    "source_note",
)

WEEKDAYS = {
    "monday": 0,
    "tuesday": 1,
    "wednesday": 2,
    "thursday": 3,
    "friday": 4,
    "saturday": 5,
    "sunday": 6,
}
ORDER_UNITS = frozenset({"PACK", "CARTON"})


@dataclass(frozen=True, slots=True)
class MasterWorkbookData:
    locations: tuple[Location, ...]
    items: tuple[Item, ...]
    item_policies: tuple[ItemPlanningPolicy, ...]
    delivery_rules: tuple[DeliveryCoverageRule, ...]
    source_version: str
    content_hash: str


@dataclass(frozen=True, slots=True)
class PlanningWorkbookData:
    forecasts: tuple[ForecastDaily, ...]
    menu_entries: tuple[MenuCalendarEntry, ...]
    bom_lines: tuple[BomLine, ...]
    source_version: str
    content_hash: str


def _content_hash(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
    except OSError as exc:
        raise InputFileError(f"cannot read {path}: {exc}") from exc
    return digest.hexdigest()


def _workbook(path: Path) -> Any:
    if load_workbook is None:
        raise InputFileError(
            "XLSX support is unavailable; install the project dependency openpyxl>=3.1,<4"
        )
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError) as exc:
        raise InputFileError(f"cannot read XLSX workbook {path}: {exc}") from exc


def _order_unit(path: Path, sheet: str, row: int, value: Any) -> str:
    parsed = _text(path, sheet, row, "order_unit", value)
    if parsed not in ORDER_UNITS:
        allowed = ", ".join(sorted(ORDER_UNITS))
        raise InputFileError(
            f"{path} sheet {sheet} row {row}, field order_unit: {parsed!r} is invalid; "
            f"use one of {allowed}"
        )
    return parsed


def _rows(path: Path, workbook: Any, sheet_name: str, headers: tuple[str, ...]) -> list[tuple[int, dict[str, Any]]]:
    if sheet_name not in workbook.sheetnames:
        raise InputFileError(
            f"{path}: missing sheet {sheet_name!r}; use the fixed V1 template without renaming tabs"
        )
    sheet = workbook[sheet_name]
    supplied = tuple(cell.value if cell.value is not None else "" for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
    supplied = tuple(str(value).strip() for value in supplied)
    if supplied != headers:
        missing = [header for header in headers if header not in supplied]
        unexpected = [header for header in supplied if header and header not in headers]
        detail: list[str] = []
        if missing:
            detail.append(f"missing {', '.join(missing)}")
        if unexpected:
            detail.append(f"unexpected {', '.join(unexpected)}")
        if not detail:
            detail.append("columns are in a different order")
        raise InputFileError(
            f"{path} sheet {sheet_name} row 1: fixed schema mismatch ({'; '.join(detail)}); "
            "copy the data into the current template and keep its row-1 headers"
        )

    rows: list[tuple[int, dict[str, Any]]] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers)), start=2):
        values = [cell.value for cell in cells]
        if not any(value not in (None, "") for value in values):
            continue
        rows.append((row_number, dict(zip(headers, values, strict=True))))
    return rows


def _text(path: Path, sheet: str, row: int, field: str, value: Any, *, optional: bool = False) -> str | None:
    if value is None:
        normalized = ""
    elif isinstance(value, bool):
        normalized = "TRUE" if value else "FALSE"
    elif isinstance(value, float) and value.is_integer():
        normalized = str(int(value))
    else:
        normalized = str(value).strip()
    if normalized:
        return normalized
    if optional:
        return None
    raise InputFileError(
        f"{path} sheet {sheet} row {row}, field {field}: value is blank; fill the highlighted template field"
    )


def _decimal(path: Path, sheet: str, row: int, field: str, value: Any, *, optional: bool = False) -> Decimal | None:
    if value in (None, ""):
        if optional:
            return None
        raise InputFileError(f"{path} sheet {sheet} row {row}, field {field}: numeric value is blank")
    try:
        return Decimal(str(value))
    except InvalidOperation as exc:
        raise InputFileError(
            f"{path} sheet {sheet} row {row}, field {field}: {value!r} is not numeric"
        ) from exc


def _integer(path: Path, sheet: str, row: int, field: str, value: Any, *, optional: bool = False) -> int | None:
    parsed = _decimal(path, sheet, row, field, value, optional=optional)
    if parsed is None:
        return None
    if parsed != parsed.to_integral_value():
        raise InputFileError(
            f"{path} sheet {sheet} row {row}, field {field}: {value!r} is not a whole number"
        )
    return int(parsed)


def _date(path: Path, sheet: str, row: int, field: str, value: Any, *, optional: bool = False) -> date | None:
    if value in (None, ""):
        if optional:
            return None
        raise InputFileError(f"{path} sheet {sheet} row {row}, field {field}: date is blank")
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError as exc:
        raise InputFileError(
            f"{path} sheet {sheet} row {row}, field {field}: {value!r} is not a real date"
        ) from exc


def _boolean(path: Path, sheet: str, row: int, field: str, value: Any) -> bool:
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"true", "1", "yes"}:
        return True
    if normalized in {"false", "0", "no"}:
        return False
    raise InputFileError(
        f"{path} sheet {sheet} row {row}, field {field}: {value!r} is not TRUE or FALSE"
    )


def _enum(path: Path, sheet: str, row: int, field: str, value: Any, enum_type: type[Any]) -> Any:
    normalized = _text(path, sheet, row, field, value)
    try:
        return enum_type(normalized)
    except ValueError as exc:
        allowed = ", ".join(member.value for member in enum_type)
        raise InputFileError(
            f"{path} sheet {sheet} row {row}, field {field}: {value!r} is invalid; use {allowed}"
        ) from exc


def _provenance(path: Path, sheet: str, row: int, value: Any) -> Provenance:
    return _enum(path, sheet, row, "provenance", value, Provenance)


def _weekday(path: Path, sheet: str, row: int, field: str, value: Any, *, optional: bool = False) -> int | None:
    normalized = _text(path, sheet, row, field, value, optional=optional)
    if normalized is None:
        return None
    weekday = WEEKDAYS.get(normalized.casefold())
    if weekday is None:
        raise InputFileError(
            f"{path} sheet {sheet} row {row}, field {field}: {normalized!r} is not an English weekday"
        )
    return weekday


def _covered_weekdays(path: Path, row: int, value: Any) -> tuple[int, ...]:
    normalized = _text(path, "Delivery_Rules", row, "covered_service_days", value, optional=True)
    if normalized is None:
        return ()
    result: list[int] = []
    for part in normalized.split(","):
        weekday = WEEKDAYS.get(part.strip().casefold())
        if weekday is None:
            raise InputFileError(
                f"{path} sheet Delivery_Rules row {row}, field covered_service_days: "
                f"{part.strip()!r} is not an English weekday"
            )
        if weekday in result:
            raise InputFileError(
                f"{path} sheet Delivery_Rules row {row}: covered service weekday {part.strip()!r} is duplicated"
            )
        result.append(weekday)
    return tuple(result)


def _supplier_id(ordering_channel: str) -> str:
    normalized = re.sub(r"[^A-Z0-9]+", "_", ordering_channel.upper()).strip("_")
    if not normalized:
        raise ValueError("ordering channel cannot produce a stable supplier ID")
    return normalized


def _unique(path: Path, sheet: str, rows: Iterable[tuple[int, Any]], key_name: str, key: Any) -> None:
    seen: dict[Any, int] = {}
    for row_number, value in rows:
        selected = key(value)
        if selected in seen:
            raise InputFileError(
                f"{path} sheet {sheet} row {row_number}, key {key_name}={selected!r}: "
                f"duplicate row; first seen at row {seen[selected]}"
            )
        seen[selected] = row_number


def load_master_template(path: Path) -> MasterWorkbookData:
    workbook = _workbook(path)
    try:
        item_rows = _rows(path, workbook, "Items", ITEM_HEADERS)
        location_rows = _rows(path, workbook, "Locations", LOCATION_HEADERS)
        delivery_rows = _rows(path, workbook, "Delivery_Rules", DELIVERY_HEADERS)

        parsed_items: list[tuple[int, Item]] = []
        policies: list[tuple[int, ItemPlanningPolicy]] = []
        for row_number, row in item_rows:
            active = _boolean(path, "Items", row_number, "active", row["active"])
            item_id = _text(path, "Items", row_number, "item_id", row["item_id"])
            storage_class = _enum(
                path, "Items", row_number, "storage_class", row["storage_class"], StorageClass
            )
            item = Item(
                item_id=item_id,
                item_name=_text(path, "Items", row_number, "item_name", row["item_name"]),
                storage_class=storage_class,
                pack_size_g=_decimal(path, "Items", row_number, "pack_size_g", row["pack_size_g"]),
                shelf_life_days=_integer(
                    path, "Items", row_number, "shelf_life_days", row["shelf_life_days"], optional=True
                ),
                min_safety_days=_decimal(
                    path, "Items", row_number, "min_safety_days", row["min_safety_days"]
                ),
                max_cover_days=_decimal(
                    path, "Items", row_number, "max_cover_days", row["max_cover_days"], optional=True
                ),
                active=active,
                provenance=Provenance.MANUAL,
            )
            ordering_channel = _text(
                path, "Items", row_number, "ordering_channel", row["ordering_channel"]
            )
            policy = ItemPlanningPolicy(
                item_id=item_id,
                item_type=_enum(path, "Items", row_number, "item_type", row["item_type"], ItemType),
                official_supplier=_text(
                    path, "Items", row_number, "official_supplier", row["official_supplier"]
                ),
                ordering_channel=ordering_channel,
                supplier_id=_supplier_id(ordering_channel),
                supplier_article_number=_text(
                    path, "Items", row_number, "po_article_no", row["po_article_no"], optional=True
                ),
                supplier_description_match=_text(
                    path,
                    "Items",
                    row_number,
                    "po_description_match",
                    row["po_description_match"],
                    optional=True,
                ),
                packs_per_order_unit=_decimal(
                    path,
                    "Items",
                    row_number,
                    "packs_per_order_unit",
                    row["packs_per_order_unit"],
                ),
                order_unit=_order_unit(path, "Items", row_number, row["order_unit"]),
                stock_qty_unit=_enum(
                    path,
                    "Items",
                    row_number,
                    "stock_qty_unit",
                    row["stock_qty_unit"],
                    StockQuantityUnit,
                ),
                apicbase_uid=_text(
                    path, "Items", row_number, "apicbase_uid", row["apicbase_uid"], optional=True
                ),
                apicbase_stock_item_name=_text(
                    path,
                    "Items",
                    row_number,
                    "apicbase_stock_item_name",
                    row["apicbase_stock_item_name"],
                    optional=True,
                ),
                lead_time_calendar_days=_integer(
                    path,
                    "Items",
                    row_number,
                    "lead_time_calendar_days",
                    row["lead_time_calendar_days"],
                ),
                shelf_life_anchor=_enum(
                    path,
                    "Items",
                    row_number,
                    "shelf_life_anchor",
                    row["shelf_life_anchor"],
                    ShelfLifeAnchor,
                )
                if row["shelf_life_anchor"] not in (None, "")
                else None,
                yield_factor=_decimal(
                    path, "Items", row_number, "yield_factor", row["yield_factor"]
                ),
                moq_order_units=_decimal(
                    path, "Items", row_number, "moq_order_units", row["moq_order_units"]
                ),
                case_multiple_order_units=_decimal(
                    path,
                    "Items",
                    row_number,
                    "case_multiple_order_units",
                    row["case_multiple_order_units"],
                ),
                data_status=_text(
                    path, "Items", row_number, "data_status", row["data_status"]
                ),
                provenance=Provenance.MANUAL,
            )
            parsed_items.append((row_number, item))
            policies.append((row_number, policy))

        parsed_locations: list[tuple[int, Location]] = []
        for row_number, row in location_rows:
            parsed_locations.append(
                (
                    row_number,
                    Location(
                        location_id=_text(
                            path, "Locations", row_number, "location_id", row["location_id"]
                        ),
                        location_name=_text(
                            path,
                            "Locations",
                            row_number,
                            "location_name",
                            row["location_name"],
                        ),
                        timezone=_text(
                            path, "Locations", row_number, "timezone", row["timezone"]
                        ),
                        active=_boolean(
                            path, "Locations", row_number, "active", row["active"]
                        ),
                    ),
                )
            )

        parsed_rules: list[tuple[int, DeliveryCoverageRule]] = []
        for row_number, row in delivery_rows:
            parsed_rules.append(
                (
                    row_number,
                    DeliveryCoverageRule(
                        delivery_rule_id=_text(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "delivery_rule_id",
                            row["delivery_rule_id"],
                        ),
                        location_id=_text(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "location_id",
                            row["location_id"],
                        ),
                        ordering_channel=_text(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "ordering_channel",
                            row["ordering_channel"],
                        ),
                        storage_class=_enum(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "storage_class",
                            row["storage_class"],
                            StorageClass,
                        ),
                        delivery_weekday=_weekday(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "delivery_weekday",
                            row["delivery_weekday"],
                            optional=True,
                        ),
                        covered_service_weekdays=_covered_weekdays(
                            path, row_number, row["covered_service_days"]
                        ),
                        order_weekday=_weekday(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "order_weekday",
                            row["order_weekday"],
                            optional=True,
                        ),
                        review_period_days=_integer(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "review_period_days",
                            row["review_period_days"],
                        ),
                        effective_from=_date(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "effective_from",
                            row["effective_from"],
                        ),
                        effective_to=_date(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "effective_to",
                            row["effective_to"],
                            optional=True,
                        ),
                        active=_boolean(
                            path, "Delivery_Rules", row_number, "active", row["active"]
                        ),
                        data_status=_text(
                            path,
                            "Delivery_Rules",
                            row_number,
                            "data_status",
                            row["data_status"],
                        ),
                        provenance=Provenance.MANUAL,
                    ),
                )
            )

        _unique(path, "Items", parsed_items, "item_id", lambda item: item.item_id)
        _unique(path, "Locations", parsed_locations, "location_id", lambda location: location.location_id)
        _unique(
            path,
            "Delivery_Rules",
            parsed_rules,
            "delivery_rule_id",
            lambda rule: rule.delivery_rule_id,
        )
        location_ids = {location.location_id for _, location in parsed_locations}
        unknown_rule_locations = sorted(
            {rule.location_id for _, rule in parsed_rules if rule.location_id not in location_ids}
        )
        if unknown_rule_locations:
            raise InputFileError(
                f"{path} Delivery_Rules references locations absent from Locations: "
                f"{', '.join(unknown_rule_locations)}"
            )
    finally:
        workbook.close()

    content_hash = _content_hash(path)
    return MasterWorkbookData(
        locations=tuple(location for _, location in parsed_locations),
        items=tuple(item for _, item in parsed_items),
        item_policies=tuple(policy for _, policy in policies),
        delivery_rules=tuple(rule for _, rule in parsed_rules),
        source_version=f"master-xlsx-{content_hash[:12]}",
        content_hash=content_hash,
    )


def load_planning_template(path: Path) -> PlanningWorkbookData:
    workbook = _workbook(path)
    try:
        demand_rows = _rows(path, workbook, "Demand_Plan", DEMAND_HEADERS)
        menu_rows = _rows(path, workbook, "Menu_Calendar", MENU_HEADERS)
        bom_rows = _rows(path, workbook, "BOM_Lines", BOM_HEADERS)

        forecasts: list[tuple[int, ForecastDaily]] = []
        for row_number, row in demand_rows:
            forecasts.append(
                (
                    row_number,
                    ForecastDaily(
                        location_id=_text(
                            path, "Demand_Plan", row_number, "location_id", row["location_id"]
                        ),
                        service_date=_date(
                            path, "Demand_Plan", row_number, "service_date", row["service_date"]
                        ),
                        dish_id=_text(
                            path, "Demand_Plan", row_number, "dish_id", row["dish_id"]
                        ),
                        forecast_portions=_decimal(
                            path,
                            "Demand_Plan",
                            row_number,
                            "forecast_portions",
                            row["forecast_portions"],
                        ),
                        forecast_version=_text(
                            path,
                            "Demand_Plan",
                            row_number,
                            "forecast_version",
                            row["forecast_version"],
                        ),
                        provenance=_provenance(
                            path, "Demand_Plan", row_number, row["provenance"]
                        ),
                    ),
                )
            )

        menu_entries: list[tuple[int, MenuCalendarEntry]] = []
        for row_number, row in menu_rows:
            menu_entries.append(
                (
                    row_number,
                    MenuCalendarEntry(
                        location_id=_text(
                            path, "Menu_Calendar", row_number, "location_id", row["location_id"]
                        ),
                        service_date=_date(
                            path, "Menu_Calendar", row_number, "service_date", row["service_date"]
                        ),
                        dish_id=_text(
                            path, "Menu_Calendar", row_number, "dish_id", row["dish_id"]
                        ),
                        menu_version=_text(
                            path, "Menu_Calendar", row_number, "menu_version", row["menu_version"]
                        ),
                        active=_boolean(
                            path, "Menu_Calendar", row_number, "active", row["active"]
                        ),
                        provenance=_provenance(
                            path, "Menu_Calendar", row_number, row["provenance"]
                        ),
                    ),
                )
            )

        bom_lines: list[tuple[int, BomLine]] = []
        for row_number, row in bom_rows:
            if not _boolean(path, "BOM_Lines", row_number, "active", row["active"]):
                continue
            bom_lines.append(
                (
                    row_number,
                    BomLine(
                        bom_line_id=_text(
                            path, "BOM_Lines", row_number, "bom_line_id", row["bom_line_id"]
                        ),
                        dish_id=_text(
                            path, "BOM_Lines", row_number, "dish_id", row["dish_id"]
                        ),
                        silo_id=_text(
                            path, "BOM_Lines", row_number, "silo_id", row["silo_id"]
                        ),
                        item_id=_text(
                            path, "BOM_Lines", row_number, "item_id", row["item_id"]
                        ),
                        grams_per_portion=_decimal(
                            path,
                            "BOM_Lines",
                            row_number,
                            "grams_per_portion",
                            row["grams_per_portion"],
                        ),
                        effective_from=_date(
                            path,
                            "BOM_Lines",
                            row_number,
                            "effective_from",
                            row["effective_from"],
                        ),
                        effective_to=_date(
                            path,
                            "BOM_Lines",
                            row_number,
                            "effective_to",
                            row["effective_to"],
                            optional=True,
                        ),
                        provenance=Provenance.MANUAL,
                    ),
                )
            )

        _unique(
            path,
            "Demand_Plan",
            forecasts,
            "location_id+dish_id+service_date+forecast_version",
            lambda row: (row.location_id, row.dish_id, row.service_date, row.forecast_version),
        )
        _unique(
            path,
            "Menu_Calendar",
            menu_entries,
            "location_id+dish_id+service_date+menu_version",
            lambda row: (row.location_id, row.dish_id, row.service_date, row.menu_version),
        )
        _unique(path, "BOM_Lines", bom_lines, "bom_line_id", lambda row: row.bom_line_id)
    finally:
        workbook.close()

    content_hash = _content_hash(path)
    return PlanningWorkbookData(
        forecasts=tuple(row for _, row in forecasts),
        menu_entries=tuple(row for _, row in menu_entries),
        bom_lines=tuple(row for _, row in bom_lines),
        source_version=f"planning-xlsx-{content_hash[:12]}",
        content_hash=content_hash,
    )
