from __future__ import annotations

import hashlib
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from types import SimpleNamespace
from typing import cast
from unittest.mock import patch

from openpyxl import Workbook  # type: ignore[import-untyped]

from apps.api.supply_planning_api.auth import AuthenticatedUser
from apps.api.supply_planning_api.config import Settings
from apps.api.supply_planning_api.repository import CanonicalStore, JsonObject, SchemaState
from apps.api.supply_planning_api.services import PlanningBackend, _date_value
from apps.api.supply_planning_api.source_rows import read_validated_sheet_rows
from apps.api.supply_planning_api.uploads import SavedUpload
from supply_planning.adapters.template_xlsx import (
    DELIVERY_HEADERS,
    ITEM_HEADERS,
    LOCATION_HEADERS,
)


class _ImportStore:
    def __init__(self) -> None:
        self.master_payload: JsonObject | None = None

    async def schema_state(self) -> SchemaState:
        return SchemaState(True)

    async def select_rows(
        self,
        table: str,
        *,
        columns: str = "*",
        filters: dict[str, str] | None = None,
        order: str | None = None,
        limit: int | None = None,
    ) -> list[JsonObject]:
        del table, columns, filters, order, limit
        return []

    async def insert_rows(self, table: str, rows: list[JsonObject]) -> list[JsonObject]:
        del table
        return rows

    async def update_rows(
        self,
        table: str,
        values: JsonObject,
        *,
        filters: dict[str, str],
    ) -> list[JsonObject]:
        del table, values, filters
        raise AssertionError("not used")

    async def delete_rows(self, table: str, *, filters: dict[str, str]) -> None:
        del table, filters
        raise AssertionError("not used")

    async def activate_master_version(
        self,
        *,
        version_id: str,
        actor_id: str,
        environment: str,
    ) -> JsonObject:
        del version_id, actor_id, environment
        raise AssertionError("not used")

    async def persist_master_import(self, payload: JsonObject) -> JsonObject:
        self.master_payload = payload
        source = dict(payload["source_import"])
        source["status"] = payload["final_status"]
        return source

    async def persist_source_import(self, payload: JsonObject) -> JsonObject:
        del payload
        raise AssertionError("not used")

    async def persist_planning_run(self, payload: JsonObject) -> str:
        del payload
        raise AssertionError("not used")


def _master_workbook(path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    items = workbook.create_sheet("Items")
    items.append(ITEM_HEADERS)
    items.append(
        (
            "ITEM_1",
            "Test item",
            "INGREDIENT",
            "Transgourmet",
            "TRANSGOURMET",
            "123",
            "Test item",
            "",
            "",
            "",
            "RT",
            1000,
            1,
            "PACK",
            "PACK",
            "UID-1",
            "Test item",
            3,
            30,
            "RECEIPT_DATE",
            2,
            1,
            14,
            1,
            1,
            True,
            "SOURCE_VALUE",
            "test",
        )
    )
    locations = workbook.create_sheet("Locations")
    locations.append(LOCATION_HEADERS)
    locations.append(
        ("LOC_1", "Demo", "Europe/Berlin", True, "SOURCE_VALUE", "test")
    )
    rules = workbook.create_sheet("Delivery_Rules")
    rules.append(DELIVERY_HEADERS)
    rules.append(
        (
            "RULE_1",
            "LOC_1",
            "TRANSGOURMET",
            "RT",
            "",
            "",
            "",
            "",
            "",
            7,
            date(2026, 8, 26),
            "",
            True,
            "SOURCE_VALUE",
            "test",
        )
    )
    workbook.save(path)


class BackendImportServiceTests(unittest.IsolatedAsyncioTestCase):
    async def test_master_workbook_becomes_one_atomic_draft_payload(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "master.xlsx"
            _master_workbook(path)
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            saved = SavedUpload(path, path.name, path.stat().st_size, digest)
            store = _ImportStore()
            backend = PlanningBackend(
                cast(CanonicalStore, store),
                Settings.model_validate(
                    {"APP_ENV": "test", "CORS_ORIGINS": "http://localhost:5173"}
                ),
            )

            result = await backend.import_master_data(
                saved,
                AuthenticatedUser(
                    "00000000-0000-0000-0000-000000000301",
                    "maintainer@example.test",
                ),
            )

        self.assertEqual("accepted", result["status"])
        assert store.master_payload is not None
        self.assertEqual("draft", store.master_payload["master_data_version"]["status"])
        self.assertEqual(1, len(store.master_payload["items"]))
        self.assertEqual(1, len(store.master_payload["item_policy_overrides"]))
        self.assertEqual(1, len(store.master_payload["locations"]))
        self.assertEqual(1, len(store.master_payload["delivery_rules"]))
        self.assertEqual(
            store.master_payload["source_import"]["id"],
            store.master_payload["master_data_version"]["source_import_id"],
        )


class ValidatedSourceRowTests(unittest.TestCase):
    def test_excel_midnight_datetime_is_normalized_to_service_date(self) -> None:
        self.assertEqual(
            date(2026, 8, 31),
            _date_value(datetime(2026, 8, 31, 0, 0)),
        )

    def test_reads_rows_when_optional_worksheet_dimension_is_absent(self) -> None:
        headers = ("item_id", "item_name")

        class DimensionlessSheet:
            @property
            def max_row(self) -> int:
                raise AssertionError("source-row reading must not require max_row")

            def cell(self, row: int, column: int) -> SimpleNamespace:
                value = headers[column - 1] if row == 1 else None
                return SimpleNamespace(value=value)

            def iter_rows(
                self,
                *,
                min_row: int,
                max_col: int,
                values_only: bool,
            ) -> tuple[tuple[object, ...], ...]:
                self.assert_iter_arguments(min_row, max_col, values_only)
                return (("ITEM_1", "Test item"),)

            @staticmethod
            def assert_iter_arguments(
                min_row: int,
                max_col: int,
                values_only: bool,
            ) -> None:
                if (min_row, max_col, values_only) != (2, len(headers), True):
                    raise AssertionError("unexpected worksheet iteration contract")

        class DimensionlessWorkbook:
            sheetnames = ["Items"]

            def __getitem__(self, name: str) -> DimensionlessSheet:
                if name != "Items":
                    raise KeyError(name)
                return DimensionlessSheet()

            def close(self) -> None:
                return None

        with patch(
            "apps.api.supply_planning_api.source_rows.load_workbook",
            return_value=DimensionlessWorkbook(),
        ):
            rows = read_validated_sheet_rows(Path("dimensionless.xlsx"), "Items", headers)

        self.assertEqual(({"item_id": "ITEM_1", "item_name": "Test item"},), rows)


if __name__ == "__main__":
    unittest.main()
