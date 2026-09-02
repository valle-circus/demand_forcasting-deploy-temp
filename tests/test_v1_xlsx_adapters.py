from __future__ import annotations

import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook

from supply_planning.adapters.apicbase_stock_xlsx import STOCK_HEADERS, normalize_apicbase_stock
from supply_planning.adapters.errors import InputFileError
from supply_planning.adapters.template_xlsx import (
    BOM_HEADERS,
    DELIVERY_HEADERS,
    DEMAND_HEADERS,
    ITEM_HEADERS,
    LOCATION_HEADERS,
    MENU_HEADERS,
    load_master_template,
    load_planning_template,
)
from supply_planning.domain.models import (
    Item,
    ItemPlanningPolicy,
    ItemType,
    ShelfLifeAnchor,
    StockQuantityUnit,
    StorageClass,
)


def _write_table(sheet, headers: tuple[str, ...], values: tuple[object, ...]) -> None:
    sheet.append(headers)
    sheet.append(values)


def _policy(item_id: str, *, uid: str | None, name: str | None) -> ItemPlanningPolicy:
    return ItemPlanningPolicy(
        item_id=item_id,
        item_type=ItemType.INGREDIENT,
        official_supplier="Transgourmet",
        ordering_channel="TRANSGOURMET",
        supplier_id="TRANSGOURMET",
        supplier_article_number="123",
        supplier_description_match="Test item",
        packs_per_order_unit=Decimal("1"),
        order_unit="PACK",
        stock_qty_unit=StockQuantityUnit.PACK,
        apicbase_uid=uid,
        apicbase_stock_item_name=name,
        lead_time_calendar_days=3,
        shelf_life_anchor=ShelfLifeAnchor.RECEIPT_DATE,
        yield_factor=Decimal("1"),
        moq_order_units=Decimal("1"),
        case_multiple_order_units=Decimal("1"),
        data_status="SOURCE_VALUE",
    )


class V1TemplateAdapterTests(unittest.TestCase):
    def test_fixed_master_and_planning_schemas_load_typed_rows(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            master_path = root / "master.xlsx"
            master = Workbook()
            master.remove(master.active)
            _write_table(
                master.create_sheet("Items"),
                ITEM_HEADERS,
                (
                    "ITEM_1", "Test item", "INGREDIENT", "Transgourmet",
                    "TRANSGOURMET", "123", "Test item", "", "", "", "RT",
                    1000, 1, "PACK", "PACK", "UID-1", "Test item", 3, 30,
                    "RECEIPT_DATE", 2, 1, 14, 1, 1, True, "SOURCE_VALUE", "test",
                ),
            )
            _write_table(
                master.create_sheet("Locations"),
                LOCATION_HEADERS,
                ("LOC_1", "Demo", "Europe/Berlin", True, "SOURCE_VALUE", "test"),
            )
            _write_table(
                master.create_sheet("Delivery_Rules"),
                DELIVERY_HEADERS,
                (
                    "RULE_1", "LOC_1", "TRANSGOURMET", "RT", "", "", "",
                    "", "", 7, date(2026, 8, 26), "", True, "SOURCE_VALUE", "test",
                ),
            )
            master.save(master_path)

            planning_path = root / "planning.xlsx"
            planning = Workbook()
            planning.remove(planning.active)
            _write_table(
                planning.create_sheet("Demand_Plan"),
                DEMAND_HEADERS,
                ("LOC_1", date(2026, 8, 31), "DISH_1", "Dish", 10, "v1", "manual", "test"),
            )
            _write_table(
                planning.create_sheet("Menu_Calendar"),
                MENU_HEADERS,
                ("LOC_1", date(2026, 8, 31), "DISH_1", "Dish", "v1", True, "manual", "test"),
            )
            _write_table(
                planning.create_sheet("BOM_Lines"),
                BOM_HEADERS,
                (
                    "BOM_1", "DISH_1", "Dish", "SILO_1", "Silo", "ITEM_1", 100,
                    date(2026, 1, 1), "", "v1", True, "SOURCE_VALUE", "test",
                ),
            )
            planning.save(planning_path)

            master_data = load_master_template(master_path)
            planning_data = load_planning_template(planning_path)

            self.assertEqual(master_data.items[0].pack_size_g, Decimal("1000"))
            self.assertEqual(master_data.delivery_rules[0].review_period_days, 7)
            self.assertEqual(planning_data.forecasts[0].forecast_portions, Decimal("10"))
            self.assertEqual(planning_data.bom_lines[0].grams_per_portion, Decimal("100"))

    def test_renamed_or_reordered_header_is_rejected_actionably(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Items"
            sheet.append(("renamed_item_id", *ITEM_HEADERS[1:]))
            workbook.create_sheet("Locations").append(LOCATION_HEADERS)
            workbook.create_sheet("Delivery_Rules").append(DELIVERY_HEADERS)
            workbook.save(path)

            with self.assertRaisesRegex(InputFileError, "fixed schema mismatch"):
                load_master_template(path)

    def test_unsupported_order_unit_is_rejected_before_persistence(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bad-order-unit.xlsx"
            workbook = Workbook()
            workbook.remove(workbook.active)
            _write_table(
                workbook.create_sheet("Items"),
                ITEM_HEADERS,
                (
                    "ITEM_1", "Test item", "INGREDIENT", "Transgourmet",
                    "TRANSGOURMET", "123", "Test item", "", "", "", "RT",
                    1000, 1, "BAG", "PACK", "UID-1", "Test item", 3, 30,
                    "RECEIPT_DATE", 2, 1, 14, 1, 1, True, "SOURCE_VALUE", "test",
                ),
            )
            _write_table(
                workbook.create_sheet("Locations"),
                LOCATION_HEADERS,
                ("LOC_1", "Demo", "Europe/Berlin", True, "SOURCE_VALUE", "test"),
            )
            _write_table(
                workbook.create_sheet("Delivery_Rules"),
                DELIVERY_HEADERS,
                (
                    "RULE_1", "LOC_1", "TRANSGOURMET", "RT", "", "", "",
                    "", "", 7, date(2026, 8, 26), "", True, "SOURCE_VALUE", "test",
                ),
            )
            workbook.save(path)

            with self.assertRaisesRegex(
                InputFileError,
                r"sheet Items row 2, field order_unit: 'BAG' is invalid; use one of CARTON, PACK",
            ):
                load_master_template(path)


class ApicbaseStockAdapterTests(unittest.TestCase):
    def test_uid_mapping_preserves_fractional_stock_and_adds_explicit_scenario_zero(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "stock.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Stock Report"
            sheet.cell(1, 1, "Stock Report for PREP-CGN")
            sheet.cell(1, 2, "Exportiert am 26.08.2026 17:10")
            headers = (
                "Stock Item Name", "UID", "Acc. Category", "Category", "Subcategory",
                "Supplier Article Numbers", "Current Stock (qty)", "Current Stock (value)",
                "Par", "Diff. with Par", "Min. Stock", "Diff with Min. Stock",
            )
            for column, header in enumerate(headers, start=1):
                sheet.cell(3, column, header)
            sheet.append(("",) * len(headers))
            for column, value in enumerate(
                ("Test item", "UID-1", "", "", "", "123", 1.25, 0, 0, 0, 0, 0),
                start=1,
            ):
                sheet.cell(4, column, value)
            workbook.save(path)

            items = (
                Item("ITEM_1", "Test item", StorageClass.RT, Decimal("1000")),
                Item("ITEM_2", "Missing item", StorageClass.RT, Decimal("500")),
            )
            result = normalize_apicbase_stock(
                path,
                location_id="LOC_1",
                timezone_name="Europe/Berlin",
                items=items,
                item_policies=(
                    _policy("ITEM_1", uid="UID-1", name="Test item"),
                    _policy("ITEM_2", uid=None, name="Missing item"),
                ),
                required_item_ids=("ITEM_1", "ITEM_2"),
                assume_zero_for_unmapped=True,
            )

            snapshots = {row.item_id: row for row in result.snapshots}
            self.assertEqual(snapshots["ITEM_1"].usable_on_hand_units, Decimal("1.25"))
            self.assertEqual(snapshots["ITEM_2"].usable_on_hand_units, Decimal("0"))
            self.assertEqual(snapshots["ITEM_2"].provenance.value, "policy_default")
            self.assertEqual(len(result.reviews), 1)
            self.assertEqual(result.counted_at.isoformat(), "2026-08-26T17:10:00+02:00")

    def test_accepts_valid_xlsx_without_optional_sheet_dimension_metadata(self) -> None:
        class DimensionlessSheet:
            @property
            def max_row(self) -> int:
                raise AssertionError("normalization must not depend on worksheet dimensions")

            def cell(self, *, row: int, column: int) -> SimpleNamespace:
                values: dict[tuple[int, int], object] = {
                    (1, 1): "Stock Report for PREP-CGN",
                    (1, 2): "Exportiert am 26.08.2026 17:10",
                    **{
                        (3, index): value
                        for index, value in enumerate(STOCK_HEADERS, start=1)
                    },
                }
                return SimpleNamespace(value=values.get((row, column)))

            def iter_rows(
                self,
                *,
                min_row: int,
                max_col: int,
                values_only: bool,
            ) -> tuple[tuple[object, ...], ...]:
                self.assert_iter_arguments(min_row, max_col, values_only)
                return (
                    (
                        "Test item",
                        "UID-1",
                        "",
                        "",
                        "",
                        "123",
                        1.25,
                        0,
                        0,
                        0,
                        0,
                        0,
                    ),
                )

            @staticmethod
            def assert_iter_arguments(
                min_row: int,
                max_col: int,
                values_only: bool,
            ) -> None:
                if (min_row, max_col, values_only) != (4, len(STOCK_HEADERS), True):
                    raise AssertionError("unexpected worksheet iteration contract")

        class DimensionlessWorkbook:
            sheetnames = ["Stock Report"]

            def __getitem__(self, name: str) -> DimensionlessSheet:
                if name != "Stock Report":
                    raise KeyError(name)
                return DimensionlessSheet()

            def close(self) -> None:
                return None

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "dimensionless.xlsx"
            path.write_bytes(b"dimensionless-fixture")
            with patch(
                "supply_planning.adapters.apicbase_stock_xlsx.load_workbook",
                return_value=DimensionlessWorkbook(),
            ):
                result = normalize_apicbase_stock(
                    path,
                    location_id="LOC_1",
                    timezone_name="Europe/Berlin",
                    items=(
                        Item("ITEM_1", "Test item", StorageClass.RT, Decimal("1000")),
                    ),
                    item_policies=(
                        _policy("ITEM_1", uid="UID-1", name="Test item"),
                    ),
                    required_item_ids=("ITEM_1",),
                    assume_zero_for_unmapped=True,
                )

        self.assertEqual(result.snapshots[0].usable_on_hand_units, Decimal("1.25"))


if __name__ == "__main__":
    unittest.main()
