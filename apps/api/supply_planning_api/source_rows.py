from __future__ import annotations

from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook  # type: ignore[import-untyped]

from supply_planning.adapters.errors import InputFileError


def read_validated_sheet_rows(
    path: Path,
    sheet_name: str,
    headers: tuple[str, ...],
) -> tuple[dict[str, Any], ...]:
    """Read display/audit fields after the authoritative template parser succeeds."""

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError) as exc:
        raise InputFileError(f"cannot reopen validated workbook {path}: {exc}") from exc
    try:
        if sheet_name not in workbook.sheetnames:
            raise InputFileError(f"{path}: missing required sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        supplied = tuple(
            "" if sheet.cell(1, column).value is None else str(sheet.cell(1, column).value).strip()
            for column in range(1, len(headers) + 1)
        )
        if supplied != headers:
            raise InputFileError(
                f"{path} sheet {sheet_name}: headers changed after validation; upload again"
            )
        rows: list[dict[str, Any]] = []
        for row_number in range(2, sheet.max_row + 1):
            values = tuple(
                sheet.cell(row_number, column).value
                for column in range(1, len(headers) + 1)
            )
            if not any(value not in (None, "") for value in values):
                continue
            rows.append(dict(zip(headers, values, strict=True)))
        return tuple(rows)
    finally:
        workbook.close()


def text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def optional_text(value: Any) -> str | None:
    normalized = text_value(value)
    return normalized or None


def temporal_text(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return text_value(value)
